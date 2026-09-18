#!/usr/bin/env python3
"""End-to-end check for per-reference internal Word links."""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skills" / "manuscript-references"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def run(*args: str) -> None:
    subprocess.run([sys.executable, *args], check=True)


def main() -> int:
    with tempfile.TemporaryDirectory() as raw_tmp:
        tmp = Path(raw_tmp)
        source = tmp / "source.docx"
        plan_path = tmp / "plan.json"
        master_path = tmp / "master.json"
        output = tmp / "output.docx"
        workbook = tmp / "references.xlsx"

        document = Document()
        document.add_paragraph("Introduction")
        document.add_paragraph("Evidence SOURCE ONE AND SOURCE TWO.")
        document.add_paragraph("Additional evidence SOURCE ONE.")
        document.save(source)

        plan = {
            "replacements": [
                {
                    "occurrence_id": "CIT-0001",
                    "location_id": "body.p000002",
                    "original_text": "SOURCE ONE AND SOURCE TWO",
                    "ref_ids": ["REF-AAAA", "REF-BBBB"],
                    "replacement_runs": [
                        {"text": "("},
                        {"text": "Smith, 2025", "ref_id": "REF-AAAA"},
                        {"text": "; "},
                        {"text": "Jones, 2026", "ref_id": "REF-BBBB"},
                        {"text": ")"},
                    ],
                    "status": "resolved",
                },
                {
                    "occurrence_id": "CIT-0002",
                    "location_id": "body.p000003",
                    "original_text": "SOURCE ONE",
                    "replacement": "(Smith, 2025)",
                    "ref_ids": ["REF-AAAA"],
                    "status": "resolved",
                },
            ],
            "bibliography": {
                "heading": "References",
                "page_break": False,
                "entries": [
                    {"ref_id": "REF-AAAA", "text": "Smith, A. (2025). First work."},
                    {"ref_id": "REF-BBBB", "text": "Jones, B. (2026). Second work."},
                ],
            },
        }
        master = {
            "references": [
                {"ref_id": "REF-AAAA", "verified_title": "First work", "status": "verified"},
                {"ref_id": "REF-BBBB", "verified_title": "Second work", "status": "verified"},
            ],
            "citation_map": [
                {
                    "occurrence_id": "CIT-0001",
                    "location_id": "body.p000002",
                    "original_text": "SOURCE ONE AND SOURCE TWO",
                    "ref_ids": ["REF-AAAA", "REF-BBBB"],
                    "formatted_citation": "(Smith, 2025; Jones, 2026)",
                    "status": "resolved",
                },
                {
                    "occurrence_id": "CIT-0002",
                    "location_id": "body.p000003",
                    "original_text": "SOURCE ONE",
                    "ref_ids": ["REF-AAAA"],
                    "formatted_citation": "(Smith, 2025)",
                    "status": "resolved",
                },
            ],
            "format_rules": {"verification_mode": "fast", "word_links": "internal"},
            "issues": [],
        }
        plan_path.write_text(json.dumps(plan), encoding="utf-8")
        master_path.write_text(json.dumps(master), encoding="utf-8")

        run(str(SKILL / "scripts" / "apply_citation_plan.py"), str(source), str(plan_path), str(output))
        run(str(SKILL / "scripts" / "build_reference_workbook.py"), str(master_path), str(workbook))
        run(str(SKILL / "scripts" / "validate_reference_package.py"), str(output), str(workbook), str(plan_path))

        with zipfile.ZipFile(output) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
        bookmarks = set(root.xpath("//w:bookmarkStart/@w:name", namespaces={"w": W_NS}))
        anchors = root.xpath("//w:hyperlink/@w:anchor", namespaces={"w": W_NS})
        assert bookmarks >= {"REF_AAAA", "REF_BBBB"}
        assert anchors.count("REF_AAAA") == 2
        assert anchors.count("REF_BBBB") == 1
        assert not root.xpath("//w:hyperlink//w:rStyle[@w:val='Hyperlink']", namespaces={"w": W_NS})

    print("internal link test passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
