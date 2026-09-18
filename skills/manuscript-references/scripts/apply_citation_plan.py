#!/usr/bin/env python3
"""Apply exact citation replacements to DOCX and append a formatted bibliography."""

from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterator

from docx import Document
from docx.document import Document as _Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Inches, Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run


def iter_block_items(parent: _Document | _Cell) -> Iterator[Paragraph | Table]:
    parent_elm = parent.element.body if isinstance(parent, _Document) else parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def location_map(document: _Document) -> dict[str, Paragraph]:
    result: dict[str, Paragraph] = {}
    body_index = 0
    table_index = 0

    def add_table(table: Table) -> None:
        nonlocal table_index
        table_index += 1
        current_table = table_index
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                para_index = 0
                for block in iter_block_items(cell):
                    if isinstance(block, Paragraph):
                        para_index += 1
                        key = f"table.t{current_table:04d}.r{row_index:04d}.c{cell_index:04d}.p{para_index:04d}"
                        result[key] = block
                    else:
                        add_table(block)

    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            body_index += 1
            result[f"body.p{body_index:06d}"] = block
        else:
            add_table(block)
    return result


def set_run_format(run: Run, spec: dict[str, Any], font_name: str = "", font_size_pt: float | None = None) -> None:
    run.bold = bool(spec.get("bold", False))
    run.italic = bool(spec.get("italic", False))
    run.underline = bool(spec.get("underline", False))
    run.font.superscript = bool(spec.get("superscript", False))
    run.font.subscript = bool(spec.get("subscript", False))
    chosen_font = str(spec.get("font_name") or font_name or "").strip()
    if chosen_font:
        run.font.name = chosen_font
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), chosen_font)
    chosen_size = spec.get("font_size_pt", font_size_pt)
    if chosen_size not in (None, ""):
        run.font.size = Pt(float(chosen_size))


def bookmark_name(ref_id: str) -> str:
    """Return a deterministic Word-safe bookmark name (40 characters max)."""
    cleaned = re.sub(r"[^A-Za-z0-9_]", "_", ref_id.strip())
    if not cleaned or not cleaned[0].isalpha():
        cleaned = "REF_" + cleaned
    return cleaned[:40]


def insert_segment_after(
    anchor_element: Any,
    style_source: Run,
    paragraph: Paragraph,
    text: str,
    spec: dict[str, Any],
    bookmark_by_ref_id: dict[str, str],
) -> tuple[Any, Run]:
    new_r = OxmlElement("w:r")
    if style_source._r.rPr is not None:
        new_r.append(copy.deepcopy(style_source._r.rPr))
    new_t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        new_t.set(qn("xml:space"), "preserve")
    new_t.text = text
    new_r.append(new_t)
    new_run = Run(new_r, paragraph)
    set_run_format(new_run, spec)

    ref_id = str(spec.get("ref_id") or "").strip()
    target = bookmark_by_ref_id.get(ref_id)
    if target:
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), target)
        hyperlink.set(qn("w:history"), "1")
        hyperlink.append(new_r)
        anchor_element.addnext(hyperlink)
        return hyperlink, new_run

    anchor_element.addnext(new_r)
    return new_r, new_run


def replacement_segments(item: dict[str, Any]) -> list[dict[str, Any]]:
    segments = item.get("replacement_runs")
    if isinstance(segments, list) and segments:
        return [{**segment, "text": str(segment.get("text", ""))} for segment in segments]
    ref_ids = [str(value) for value in (item.get("ref_ids") or []) if str(value).strip()]
    segment: dict[str, Any] = {"text": str(item.get("replacement", ""))}
    if len(ref_ids) == 1:
        segment["ref_id"] = ref_ids[0]
    return [segment]


def replace_exact(
    paragraph: Paragraph,
    original: str,
    segments: list[dict[str, Any]],
    bookmark_by_ref_id: dict[str, str],
    occurrence_index: int | None = None,
) -> str:
    full_text = "".join(run.text for run in paragraph.runs)
    positions: list[int] = []
    start = 0
    while True:
        found = full_text.find(original, start)
        if found < 0:
            break
        positions.append(found)
        start = found + max(1, len(original))
    if not positions:
        raise ValueError("original text not found at location")
    if occurrence_index is None:
        if len(positions) != 1:
            raise ValueError(f"original text occurs {len(positions)} times; occurrence_index is required")
        match_start = positions[0]
    else:
        if occurrence_index < 1 or occurrence_index > len(positions):
            raise ValueError("occurrence_index is outside the available matches")
        match_start = positions[occurrence_index - 1]
    match_end = match_start + len(original)

    run_ranges: list[tuple[int, int, Run]] = []
    cursor = 0
    for run in paragraph.runs:
        run_ranges.append((cursor, cursor + len(run.text), run))
        cursor += len(run.text)

    start_info = next((item for item in run_ranges if item[0] <= match_start < item[1] or (match_start == item[1] and item[0] == item[1])), None)
    end_info = next((item for item in run_ranges if item[0] < match_end <= item[1]), None)
    if start_info is None or end_info is None:
        raise ValueError("match could not be mapped to Word runs")

    start_index = run_ranges.index(start_info)
    end_index = run_ranges.index(end_info)
    start_begin, _, start_run = start_info
    end_begin, _, end_run = end_info
    prefix = start_run.text[: match_start - start_begin]
    suffix = end_run.text[match_end - end_begin :]

    start_run.text = prefix
    for index in range(start_index + 1, end_index):
        run_ranges[index][2].text = ""
    if end_index != start_index:
        end_run.text = suffix

    anchor_element = start_run._r
    for segment in segments:
        anchor_element, _ = insert_segment_after(
            anchor_element,
            start_run,
            paragraph,
            str(segment.get("text", "")),
            segment,
            bookmark_by_ref_id,
        )
    if end_index == start_index and suffix:
        insert_segment_after(anchor_element, start_run, paragraph, suffix, {}, bookmark_by_ref_id)

    return "".join(str(segment.get("text", "")) for segment in segments)


def add_bibliography(document: _Document, spec: dict[str, Any]) -> tuple[int, dict[str, str]]:
    entries = spec.get("entries") or []
    if not entries:
        return 0, {}
    if spec.get("page_break", True):
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    heading = document.add_paragraph()
    try:
        heading.style = str(spec.get("heading_style") or "Heading 1")
    except KeyError:
        pass
    heading_run = heading.add_run(str(spec.get("heading") or "References"))
    if spec.get("heading_bold") is not None:
        heading_run.bold = bool(spec.get("heading_bold"))
    font_name = str(spec.get("font_name") or "")
    font_size = spec.get("font_size_pt")
    if font_name:
        heading_run.font.name = font_name
        heading_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    if font_size not in (None, ""):
        heading_run.font.size = Pt(float(font_size))

    hanging = float(spec.get("hanging_indent_in", 0.5))
    line_spacing = float(spec.get("line_spacing", 1.0))
    space_after = float(spec.get("space_after_pt", 0))
    bookmark_by_ref_id: dict[str, str] = {}
    used_bookmarks: set[str] = set()
    bookmark_id = 1
    for entry in entries:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(hanging)
        paragraph.paragraph_format.first_line_indent = Inches(-hanging)
        paragraph.paragraph_format.line_spacing = line_spacing
        paragraph.paragraph_format.space_after = Pt(space_after)
        runs = entry.get("runs")
        if not isinstance(runs, list) or not runs:
            runs = [{"text": str(entry.get("text", ""))}]
        for run_spec in runs:
            run = paragraph.add_run(str(run_spec.get("text", "")))
            set_run_format(run, run_spec, font_name, float(font_size) if font_size not in (None, "") else None)
        ref_id = str(entry.get("ref_id") or "").strip()
        if ref_id:
            name = bookmark_name(ref_id)
            if name in used_bookmarks:
                raise ValueError(f"duplicate bibliography bookmark: {name}")
            used_bookmarks.add(name)
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(bookmark_id))
            start.set(qn("w:name"), name)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(bookmark_id))
            paragraph._p.insert(0, start)
            paragraph._p.append(end)
            bookmark_by_ref_id[ref_id] = name
            bookmark_id += 1
    return len(entries), bookmark_by_ref_id


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("plan_json", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("--audit-json", type=Path)
    args = parser.parse_args()

    try:
        document = Document(args.input_docx)
        plan = json.loads(args.plan_json.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError, Exception) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2

    locations = location_map(document)
    try:
        bibliography_count, bookmark_by_ref_id = add_bibliography(document, plan.get("bibliography") or {})
    except ValueError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    results: list[dict[str, Any]] = []
    failures = 0
    applied = 0
    for item in plan.get("replacements") or []:
        result = {
            "occurrence_id": item.get("occurrence_id", ""),
            "location_id": item.get("location_id", ""),
            "status": "skipped",
            "message": "",
        }
        if item.get("status") != "resolved":
            result["message"] = "replacement status is not resolved"
            results.append(result)
            continue
        paragraph = locations.get(str(item.get("location_id", "")))
        if paragraph is None:
            failures += 1
            result["status"] = "failed"
            result["message"] = "location_id not found"
            results.append(result)
            continue
        try:
            rendered = replace_exact(
                paragraph,
                str(item.get("original_text", "")),
                replacement_segments(item),
                bookmark_by_ref_id,
                int(item["occurrence_index"]) if item.get("occurrence_index") not in (None, "") else None,
            )
            applied += 1
            result["status"] = "applied"
            result["rendered_citation"] = rendered
        except ValueError as exc:
            failures += 1
            result["status"] = "failed"
            result["message"] = str(exc)
        results.append(result)

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)

    audit = {
        "input_file": args.input_docx.name,
        "output_file": args.output_docx.name,
        "replacement_count": len(plan.get("replacements") or []),
        "applied_count": applied,
        "failure_count": failures,
        "bibliography_count": bibliography_count,
        "results": results,
    }
    if args.audit_json:
        args.audit_json.parent.mkdir(parents=True, exist_ok=True)
        args.audit_json.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"applied {applied} replacement(s); failures={failures}; bibliography entries={bibliography_count}")
    return 3 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
