#!/usr/bin/env python3
"""Build the reusable manuscript-reference workbook from master JSON."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REFERENCE_COLUMNS = [
    ("Ref_ID", "ref_id"),
    ("Type", "type"),
    ("Original_Reference", "original_reference"),
    ("Original_Title", "original_title"),
    ("Verified_Title", "verified_title"),
    ("Authors", "authors"),
    ("Year", "year"),
    ("Source_Title", "source_title"),
    ("Volume", "volume"),
    ("Issue", "issue"),
    ("Pages_or_Article_Number", "pages_or_article_number"),
    ("DOI", "doi"),
    ("URL", "url"),
    ("Verification_Source", "verification_source"),
    ("Status", "status"),
    ("Manual_Review", "manual_review"),
    ("Notes", "notes"),
]

CITATION_COLUMNS = [
    ("Occurrence_ID", "occurrence_id"),
    ("Location_ID", "location_id"),
    ("Section", "section"),
    ("Context", "context"),
    ("Original_Text", "original_text"),
    ("Ref_IDs", "ref_ids"),
    ("Formatted_Citation", "formatted_citation"),
    ("Status", "status"),
    ("Notes", "notes"),
]

ISSUE_COLUMNS = [
    ("Issue_ID", "issue_id"),
    ("Location_ID", "location_id"),
    ("Original_Text", "original_text"),
    ("Issue_Type", "issue_type"),
    ("Description", "description"),
    ("Required_Action", "required_action"),
    ("Status", "status"),
]

HEADER_FILL = PatternFill("solid", fgColor="3559A6")
HEADER_FONT = Font(color="FFFFFF", bold=True)
STATUS_FILLS = {
    "verified": PatternFill("solid", fgColor="D9EAD3"),
    "user_approved": PatternFill("solid", fgColor="D9EAD3"),
    "resolved": PatternFill("solid", fgColor="D9EAD3"),
    "ambiguous": PatternFill("solid", fgColor="FFF2CC"),
    "unresolved": PatternFill("solid", fgColor="F4CCCC"),
    "open": PatternFill("solid", fgColor="F4CCCC"),
    "closed": PatternFill("solid", fgColor="D9EAD3"),
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(clean(item) for item in value)
    return re.sub(r"\s+", " ", str(value)).strip()


def excel_safe(value: Any) -> str:
    text = clean(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKC", clean(value)).casefold()
    return " ".join(re.sub(r"[\W_]+", " ", text).split())


def stable_ref_id(record: dict[str, Any]) -> str:
    current = clean(record.get("ref_id"))
    if current:
        return current
    doi = clean(record.get("doi")).lower()
    doi = re.sub(r"^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)", "", doi)
    identity_text = (
        record.get("verified_title")
        or record.get("original_title")
        or record.get("original_reference")
        or json.dumps(record, ensure_ascii=False, sort_keys=True)
    )
    key = f"doi:{doi}" if doi else f"title:{normalize(identity_text)}"
    digest = hashlib.sha1(key.encode("utf-8")).hexdigest()[:10].upper()
    return f"REF-{digest}"


def style_sheet(ws, widths: dict[str, float]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            value = clean(cell.value).lower()
            if value in STATUS_FILLS:
                cell.fill = STATUS_FILLS[value]
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def write_table(ws, columns: list[tuple[str, str]], records: list[dict[str, Any]]) -> None:
    ws.append([heading for heading, _ in columns])
    for record in records:
        ws.append([excel_safe(record.get(key)) for _, key in columns])


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("master_json", type=Path)
    parser.add_argument("output_xlsx", type=Path)
    args = parser.parse_args()

    try:
        data = json.loads(args.master_json.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2

    references = data.get("references") or []
    citation_map = data.get("citation_map") or []
    issues = data.get("issues") or []
    format_rules = data.get("format_rules") or {}
    if not isinstance(references, list) or not isinstance(citation_map, list) or not isinstance(issues, list):
        print("error: references, citation_map, and issues must be arrays", file=sys.stderr)
        return 2

    seen_ids: set[str] = set()
    for record in references:
        record["ref_id"] = stable_ref_id(record)
        if record["ref_id"] in seen_ids:
            print(f"error: duplicate Ref_ID {record['ref_id']}", file=sys.stderr)
            return 2
        seen_ids.add(record["ref_id"])

    workbook = Workbook()
    references_ws = workbook.active
    references_ws.title = "References"
    write_table(references_ws, REFERENCE_COLUMNS, references)
    style_sheet(
        references_ws,
        {"A": 18, "B": 18, "C": 55, "D": 40, "E": 45, "F": 38, "G": 10, "H": 30,
         "I": 10, "J": 10, "K": 24, "L": 28, "M": 42, "N": 42, "O": 16, "P": 16, "Q": 35},
    )

    citation_ws = workbook.create_sheet("Citation_Map")
    write_table(citation_ws, CITATION_COLUMNS, citation_map)
    style_sheet(citation_ws, {"A": 16, "B": 28, "C": 24, "D": 55, "E": 55, "F": 32, "G": 35, "H": 16, "I": 35})

    rules_ws = workbook.create_sheet("Format_Rules")
    rules_ws.append(["Rule", "Applied_Value"])
    if isinstance(format_rules, dict):
        for key, value in format_rules.items():
            rules_ws.append([excel_safe(key), excel_safe(value)])
    else:
        rules_ws.append(["requested_style", excel_safe(format_rules)])
    style_sheet(rules_ws, {"A": 35, "B": 95})

    issues_ws = workbook.create_sheet("Issues")
    write_table(issues_ws, ISSUE_COLUMNS, issues)
    style_sheet(issues_ws, {"A": 16, "B": 28, "C": 60, "D": 20, "E": 50, "F": 35, "G": 14})

    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 30

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_xlsx)
    print(f"created workbook with {len(references)} reference(s), {len(citation_map)} occurrence(s), and {len(issues)} issue(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
