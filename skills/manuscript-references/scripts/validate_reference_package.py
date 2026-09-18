#!/usr/bin/env python3
"""Validate consistency among the revised DOCX, workbook, and citation plan."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as _Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def iter_block_items(parent: _Document | _Cell) -> Iterator[Paragraph | Table]:
    parent_elm = parent.element.body if isinstance(parent, _Document) else parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def document_text(document: _Document) -> str:
    texts: list[str] = []

    def visit(parent: _Document | _Cell) -> None:
        for block in iter_block_items(parent):
            if isinstance(block, Paragraph):
                texts.append(block.text)
            else:
                for row in block.rows:
                    for cell in row.cells:
                        visit(cell)

    visit(document)
    return "\n".join(texts)


def rendered_replacement(item: dict) -> str:
    runs = item.get("replacement_runs")
    if isinstance(runs, list) and runs:
        return "".join(str(run.get("text", "")) for run in runs)
    return str(item.get("replacement", ""))


def rendered_entry(entry: dict) -> str:
    runs = entry.get("runs")
    if isinstance(runs, list) and runs:
        return "".join(str(run.get("text", "")) for run in runs)
    return str(entry.get("text", ""))


def data_row_count(ws) -> int:
    return sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in row))


def internal_link_inventory(docx_path: Path) -> tuple[set[str], Counter[str]]:
    with zipfile.ZipFile(docx_path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    bookmarks = {
        str(value)
        for value in root.xpath("//w:bookmarkStart/@w:name", namespaces={"w": W_NS})
        if value
    }
    anchors = Counter(
        str(value)
        for value in root.xpath("//w:hyperlink/@w:anchor", namespaces={"w": W_NS})
        if value
    )
    return bookmarks, anchors


def bookmark_name(ref_id: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_]", "_", ref_id.strip())
    if not cleaned or not cleaned[0].isalpha():
        cleaned = "REF_" + cleaned
    return cleaned[:40]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("revised_docx", type=Path)
    parser.add_argument("reference_workbook", type=Path)
    parser.add_argument("plan_json", type=Path)
    parser.add_argument("--report-json", type=Path)
    args = parser.parse_args()

    try:
        document = Document(args.revised_docx)
        workbook = load_workbook(args.reference_workbook, read_only=True, data_only=True)
        plan = json.loads(args.plan_json.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2

    errors: list[str] = []
    warnings: list[str] = []
    try:
        bookmarks, anchors = internal_link_inventory(args.revised_docx)
    except Exception as exc:
        errors.append(f"could not inspect internal links: {exc}")
        bookmarks, anchors = set(), Counter()
    required_sheets = {"References", "Citation_Map", "Format_Rules", "Issues"}
    missing_sheets = sorted(required_sheets - set(workbook.sheetnames))
    if missing_sheets:
        errors.append("missing workbook sheets: " + ", ".join(missing_sheets))

    text = document_text(document)
    expected_citations = Counter(
        rendered_replacement(item)
        for item in plan.get("replacements") or []
        if item.get("status") == "resolved" and rendered_replacement(item)
    )
    for citation, expected_count in expected_citations.items():
        actual_count = text.count(citation)
        if actual_count < expected_count:
            errors.append(f"citation occurs {actual_count} time(s), expected at least {expected_count}: {citation}")

    bibliography_entries = plan.get("bibliography", {}).get("entries") or []
    bibliography_ref_ids = {
        str(entry.get("ref_id") or "").strip()
        for entry in bibliography_entries
        if str(entry.get("ref_id") or "").strip()
    }
    for entry in bibliography_entries:
        rendered = rendered_entry(entry)
        if rendered and rendered not in text:
            errors.append(f"bibliography entry not found in document: {rendered[:100]}")
        ref_id = str(entry.get("ref_id") or "").strip()
        if not ref_id:
            errors.append(f"bibliography entry lacks Ref_ID: {rendered[:100]}")
        elif bookmark_name(ref_id) not in bookmarks:
            errors.append(f"bibliography bookmark not found for {ref_id}")

    expected_link_counts: Counter[str] = Counter()
    for item in plan.get("replacements") or []:
        if item.get("status") != "resolved":
            continue
        ref_ids = [str(value) for value in (item.get("ref_ids") or []) if str(value).strip()]
        segments = item.get("replacement_runs")
        linked_ref_ids = {
            str(segment.get("ref_id") or "").strip()
            for segment in segments
            if str(segment.get("ref_id") or "").strip()
        } if isinstance(segments, list) and segments else (set(ref_ids) if len(ref_ids) == 1 else set())
        missing_links = sorted(set(ref_ids) - linked_ref_ids)
        extra_links = sorted(linked_ref_ids - set(ref_ids))
        if missing_links:
            errors.append(
                f"citation {item.get('occurrence_id', '')} lacks per-reference link segments: {', '.join(missing_links)}"
            )
        if extra_links:
            errors.append(
                f"citation {item.get('occurrence_id', '')} links Ref_IDs absent from ref_ids: {', '.join(extra_links)}"
            )
        for ref_id in linked_ref_ids:
            expected_link_counts[bookmark_name(ref_id)] += 1
            if ref_id not in bibliography_ref_ids:
                errors.append(f"citation links to missing bibliography Ref_ID: {ref_id}")
    for target, expected_count in expected_link_counts.items():
        if anchors[target] < expected_count:
            errors.append(
                f"internal hyperlinks to {target}={anchors[target]}, expected at least {expected_count}"
            )

    if "Citation_Map" in workbook.sheetnames:
        workbook_occurrences = data_row_count(workbook["Citation_Map"])
        plan_occurrences = len(plan.get("replacements") or [])
        if workbook_occurrences != plan_occurrences:
            errors.append(f"Citation_Map rows={workbook_occurrences}, plan replacements={plan_occurrences}")
    if "References" in workbook.sheetnames:
        workbook_references = data_row_count(workbook["References"])
        bibliography_count = len(bibliography_entries)
        if workbook_references < bibliography_count:
            errors.append(f"References rows={workbook_references}, bibliography entries={bibliography_count}")

    unresolved_items = [item for item in plan.get("replacements") or [] if item.get("status") != "resolved"]
    for item in unresolved_items:
        original = str(item.get("original_text", ""))
        if original and original not in text:
            errors.append(f"unresolved original text is missing: {original[:100]}")

    if not bibliography_entries:
        warnings.append("bibliography has no entries")

    report = {
        "status": "passed" if not errors else "failed",
        "errors": errors,
        "warnings": warnings,
        "resolved_occurrences": sum(expected_citations.values()),
        "bibliography_entries": len(bibliography_entries),
        "internal_hyperlinks": sum(anchors.values()),
        "workbook_sheets": workbook.sheetnames,
    }
    if args.report_json:
        args.report_json.parent.mkdir(parents=True, exist_ok=True)
        args.report_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    return 3 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
