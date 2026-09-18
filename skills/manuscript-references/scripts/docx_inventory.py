#!/usr/bin/env python3
"""Extract likely reference-bearing DOCX paragraphs into compact JSON."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as _Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P


YEAR_RE = re.compile(r"(?:\(|\b)(?:18|19|20)\d{2}[a-z]?(?:\)|\b)")
DOI_RE = re.compile(r"(?:https?://(?:dx\.)?doi\.org/|\bdoi\s*:\s*|\b10\.\d{4,9}/)\S+", re.I)
URL_RE = re.compile(r"https?://\S+", re.I)
VOLUME_RE = re.compile(r"\b\d{1,4}\s*(?:\(\s*\d{1,4}\s*\))?\s*[,;:]\s*(?:[A-Za-z]?\d{2,}|\d+\s*[-–]\s*\d+)\b")
AUTHOR_RE = re.compile(r"\b[A-Z][A-Za-z'’\-]+\s*,\s*(?:[A-Z]\.?\s*){1,4}")
JOURNAL_RE = re.compile(r"\b(?:journal|transactions|proceedings|review|science|transportation|research|press|publisher|report|thesis|dissertation)\b", re.I)


def iter_block_items(parent: _Document | _Cell) -> Iterator[Paragraph | Table]:
    parent_elm = parent.element.body if isinstance(parent, _Document) else parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def flatten_paragraphs(document: _Document) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    body_index = 0
    table_index = 0

    def add_table(table: Table, prefix: str) -> None:
        nonlocal table_index
        table_index += 1
        current_table = table_index
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                para_index = 0
                for block in iter_block_items(cell):
                    if isinstance(block, Paragraph):
                        para_index += 1
                        rows.append(
                            {
                                "location_id": f"table.t{current_table:04d}.r{row_index:04d}.c{cell_index:04d}.p{para_index:04d}",
                                "text": block.text,
                                "style": block.style.name if block.style else "",
                                "container": prefix,
                            }
                        )
                    else:
                        add_table(block, f"nested:{prefix}")

    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            body_index += 1
            rows.append(
                {
                    "location_id": f"body.p{body_index:06d}",
                    "text": block.text,
                    "style": block.style.name if block.style else "",
                    "container": "body",
                }
            )
        else:
            add_table(block, "table")
    return rows


def reference_score(text: str) -> tuple[int, list[str]]:
    reasons: list[str] = []
    score = 0
    if DOI_RE.search(text):
        score += 4
        reasons.append("doi")
    if URL_RE.search(text):
        score += 2
        reasons.append("url")
    years = YEAR_RE.findall(text)
    if years:
        score += min(2, len(years))
        reasons.append("year")
    if VOLUME_RE.search(text):
        score += 2
        reasons.append("volume_pages")
    if AUTHOR_RE.search(text) or re.search(r"\bet\s+al\.\b", text, re.I):
        score += 2
        reasons.append("author_pattern")
    if JOURNAL_RE.search(text):
        score += 1
        reasons.append("source_terms")
    if text.count(",") >= 4 or text.count(";") >= 2:
        score += 1
        reasons.append("bibliographic_punctuation")
    return score, reasons


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_json", type=Path)
    parser.add_argument("--threshold", type=int, default=3)
    parser.add_argument("--include-all", action="store_true")
    args = parser.parse_args()

    try:
        document = Document(args.input_docx)
    except Exception as exc:
        print(f"error: cannot open DOCX: {exc}", file=sys.stderr)
        return 2

    paragraphs = flatten_paragraphs(document)
    candidates: list[dict[str, object]] = []
    current_section = ""
    for index, row in enumerate(paragraphs):
        text = row["text"].strip()
        style = row["style"].lower()
        if text and style.startswith("heading"):
            current_section = text
        score, reasons = reference_score(text)
        if args.include_all or (text and score >= args.threshold):
            previous_text = paragraphs[index - 1]["text"] if index > 0 else ""
            next_text = paragraphs[index + 1]["text"] if index + 1 < len(paragraphs) else ""
            candidates.append(
                {
                    **row,
                    "section": current_section,
                    "score": score,
                    "reasons": reasons,
                    "previous_text": previous_text,
                    "next_text": next_text,
                }
            )

    payload = {
        "source_file": args.input_docx.name,
        "paragraph_count": len(paragraphs),
        "candidate_count": len(candidates),
        "threshold": args.threshold,
        "candidates": candidates,
    }
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"extracted {len(candidates)} candidate paragraph(s) from {len(paragraphs)} paragraph(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
